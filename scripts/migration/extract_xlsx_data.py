#!/usr/bin/env python3
"""Extract customer/order rows from one or more XLSX files into canonical raw CSV."""

from __future__ import annotations

import argparse
import csv
import hashlib
import re
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path
from typing import Any

from common import parse_source_order_id

CANONICAL_HEADERS = [
    "name",
    "address",
    "pincode",
    "phone",
    "phone_display",
    "payment_type",
    "amount",
    "order_id",
    "tracking_id",
    "items",
    "notes",
    "_id",
    "created_ts",
    "printed_count",
]

COLUMN_HINTS: dict[str, list[str]] = {
    "name": ["name", "customer", "customername", "recipient"],
    "address": ["address", "deliveryaddress", "fulladdress", "addr"],
    "pincode": ["pincode", "pin", "postal", "zipcode", "zip"],
    "phone": ["phone", "mobile", "contact", "whatsapp", "number"],
    "phone_display": ["phone_display", "phonedisplay", "displayphone"],
    "payment_type": ["paymenttype", "payment", "paymode", "modeofpayment", "paid"],
    "amount": ["amount", "total", "totalamount", "ordertotal", "grandtotal", "cost"],
    "order_id": ["orderid", "order_id", "sourceorderid", "orderno", "ordernumber"],
    "tracking_id": ["trackingid", "awb", "consignment"],
    "items": ["items", "item", "products", "product", "description"],
    "notes": ["notes", "note", "remark", "remarks", "comment", "comments"],
    "_id": ["_id", "sourceid", "rowid"],
    "created_ts": ["createdts", "created", "createdat", "timestamp", "date", "orderdate"],
    "printed_count": ["printedcount", "printed", "printcount"],
}

# Sheets that contain metadata/leads but not order rows — skip entirely
NON_ORDER_SHEETS: set[str] = {
    "order",
    "color coding",
    "customer call",
    "customer call1",
    "leads",
    "sheet8",
    "sheet9",
}

ITEM_COLUMN_HINTS = ["item", "items", "product", "products", "description"]

KNOWN_ITEM_ALIASES = {
    "ഏലക്ക": "Cardamom",
    "എലക്ക": "Cardamom",
    "cardamom": "Cardamom",
    "ഗ്രാമ്പൂ": "Clove",
    "ഗ്രാമ്പു": "Clove",
    "clove": "Clove",
    "കുരുമുളക്": "Black Pepper",
    "കുരുമുളക": "Black Pepper",
    "black pepper": "Black Pepper",
    "pepper": "Black Pepper",
    "cinnamon": "Ceylon True Cinnamon",
}


def parse_args() -> argparse.Namespace:
    p = argparse.ArgumentParser(description="Extract migration rows from XLSX files")
    p.add_argument(
        "--input",
        action="append",
        required=True,
        help="Input XLSX file path (use multiple --input for multiple files)",
    )
    p.add_argument("--output", required=True, help="Output raw CSV path")
    p.add_argument(
        "--unknown-items-output",
        help="Optional CSV path with unknown item names that need translation",
    )
    p.add_argument("--batch-tag", default="xlsx-batch", help="Batch tag added to notes")
    return p.parse_args()


def norm_header(value: str | None) -> str:
    text = str(value or "").strip().lower()
    text = re.sub(r"[^a-z0-9]+", "", text)
    return text


def as_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, date):
        return datetime.combine(value, datetime.min.time()).isoformat()
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return f"{value:.6f}".rstrip("0").rstrip(".")
    return str(value).strip()


def format_phone_display(phone: str) -> str:
    digits = re.sub(r"\D", "", phone)
    if len(digits) == 12 and digits.startswith("91"):
        digits = digits[2:]
    if len(digits) == 10:
        return f"+91 {digits[:5]} {digits[5:]}"
    return phone


def parse_amount(value: str) -> float | None:
    raw = (value or "").strip()
    if not raw:
        return None
    # Handle expressions like "340-60\n=280" — take the final value after "="
    eq_match = re.search(r"=\s*(\d+(?:\.\d+)?)\s*$", raw.replace("\n", " "))
    if eq_match:
        try:
            return float(eq_match.group(1))
        except ValueError:
            pass
    cleaned = re.sub(r"[^0-9.]", "", raw.split("\n")[0])  # take first line only
    if not cleaned:
        return None
    try:
        return float(cleaned)
    except ValueError:
        return None


def format_amount(value: str) -> str:
    amt = parse_amount(value)
    if amt is None:
        return (value or "").strip()
    if float(amt).is_integer():
        return str(int(amt))
    return f"{amt:.2f}".rstrip("0").rstrip(".")


def normalize_order_id(value: str) -> str:
    text = (value or "").strip()
    if not text:
        return ""

    # "Indian Post" / "Indian Post Ayakanam" → no trackable ID
    if re.match(r"(?i)^indian\s+post", text):
        return ""

    # Combined IDs like "25090323 & IN_25092693" → take the first part
    if "&" in text:
        text = text.split("&")[0].strip()

    text = re.sub(r"\s+", "", text)
    if re.match(r"(?i)^PUR($|[_/\-]|\d)", text):
        return "PUUR" + text[3:]
    return text


def detect_columns(headers: list[str]) -> tuple[dict[str, int], list[int], int]:
    mapping: dict[str, int] = {}
    normalized = [norm_header(h) for h in headers]

    for idx, cell in enumerate(normalized):
        if not cell:
            continue
        for canon, hints in COLUMN_HINTS.items():
            if canon in mapping:
                continue
            if any(hint in cell for hint in hints):
                mapping[canon] = idx

    item_cols: list[int] = []
    for idx, cell in enumerate(normalized):
        if cell and any(h in cell for h in ITEM_COLUMN_HINTS):
            item_cols.append(idx)

    if not item_cols and "items" in mapping:
        item_cols = [mapping["items"]]

    score = len(mapping)
    if "order_id" in mapping:
        score += 3
    if "amount" in mapping:
        score += 2
    if "name" in mapping:
        score += 2
    if "phone" in mapping:
        score += 2
    if item_cols:
        score += 2
    return mapping, sorted(set(item_cols)), score


def infer_items_column(ws: Any, header_row: int, col_map: dict[str, int]) -> list[int]:
    """Fallback items-column detector for sheets where the header is a template string.

    Scans the first 10 data rows looking for an unmapped column whose values
    contain alphabetic (or Malayalam) text that looks like product descriptions.
    """
    mapped_cols = set(col_map.values())
    scores: dict[int, int] = {}
    product_kws = [
        "elakka", "pepper", "peppar", "cardamom", "cardamon",
        "clove", "gm", "gram", "kg", "cinnamon",
        "ഏലക്ക", "കുരുമുളക്", "ഗ്രാമ്പൂ", "കറുവ", "കുരുമുളക",
    ]
    rows_checked = 0
    for row in ws.iter_rows(min_row=header_row + 1, max_row=header_row + 10, values_only=True):
        for idx, cell_val in enumerate(row):
            text = as_text(cell_val).strip()
            if not text or idx in mapped_cols:
                continue
            if re.search(r"[a-zA-Z\u0D00-\u0D7F]", text) and len(text) > 3:
                scores[idx] = scores.get(idx, 0) + 1
                if any(kw.lower() in text.lower() for kw in product_kws):
                    scores[idx] = scores.get(idx, 0) + 5
        rows_checked += 1
        if rows_checked >= 10:
            break

    if not scores:
        return []
    best = max(scores, key=lambda k: scores[k])
    return [best] if scores[best] >= 3 else []


def find_header_row(ws: Any, max_scan_rows: int = 30) -> tuple[int, dict[str, int], list[int]]:
    best_row = 1
    best_map: dict[str, int] = {}
    best_items: list[int] = []
    best_score = -1

    for row_idx, row in enumerate(
        ws.iter_rows(min_row=1, max_row=max_scan_rows, values_only=True),
        start=1,
    ):
        headers = [as_text(v) for v in row]
        mapping, item_cols, score = detect_columns(headers)
        if score > best_score:
            best_row = row_idx
            best_map = mapping
            best_items = item_cols
            best_score = score

    if best_score < 5:
        raise ValueError(f"Could not detect header row for sheet '{ws.title}'")
    return best_row, best_map, best_items


def map_item_name(token: str) -> tuple[str, bool]:
    raw = token.strip()
    lower = raw.lower()

    for key, target in KNOWN_ITEM_ALIASES.items():
        k = key.lower()
        if k == lower or k in lower:
            return target, True

    return raw, False


def split_item_tokens(raw_items: str) -> list[str]:
    if not raw_items.strip():
        return []

    parts = re.split(r"[+\n,;/]|\band\b|&", raw_items, flags=re.IGNORECASE)
    tokens: list[str] = []
    for part in parts:
        t = re.sub(r"^\s*\d+[\)\.\-:]\s*", "", part.strip())
        if t:
            tokens.append(t)
    return tokens


def format_price(value: float) -> str:
    if float(value).is_integer():
        return str(int(value))
    return f"{value:.2f}".rstrip("0").rstrip(".")


def normalize_items(raw_items: str, amount_text: str) -> tuple[str, list[str]]:
    text = (raw_items or "").strip()
    if not text:
        return "", []

    # If rows already contain parseable quantity/price lines, keep as-is.
    if re.search(r"\bx\s*Rs\b", text, flags=re.IGNORECASE):
        return text, []

    tokens = split_item_tokens(text)
    if not tokens:
        return text, []

    mapped_tokens: list[str] = []
    unknown_tokens: list[str] = []
    for token in tokens:
        mapped, known = map_item_name(token)
        mapped_tokens.append(mapped)
        if not known:
            unknown_tokens.append(token)

    amount = parse_amount(amount_text)
    if amount is None:
        lines = [f"{name} - 1 x Rs 0" for name in mapped_tokens]
        return "\n".join(lines), unknown_tokens

    count = max(1, len(mapped_tokens))
    each = round(amount / count, 2)
    prices = [each] * count
    if count > 1:
        prices[-1] = round(amount - sum(prices[:-1]), 2)

    lines = [f"{name} - 1 x Rs {format_price(price)}" for name, price in zip(mapped_tokens, prices)]
    return "\n".join(lines), unknown_tokens


def to_iso_created_ts(raw: str, order_id: str, sheet_fallback_date: str = "") -> str:
    text = (raw or "").strip()
    if text:
        normalized = text.replace("Z", "+00:00")
        try:
            return datetime.fromisoformat(normalized).isoformat()
        except ValueError:
            pass

        for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%Y/%m/%d"):
            try:
                return datetime.strptime(text, fmt).isoformat()
            except ValueError:
                continue

    parsed = parse_source_order_id(order_id)
    if parsed:
        return f"{parsed.iso_date}T00:00:00"

    if sheet_fallback_date:
        return f"{sheet_fallback_date}T00:00:00"
    return ""


_MONTH_MAP: dict[str, str] = {
    "jan": "01", "january": "01",
    "feb": "02", "february": "02",
    "mar": "03", "march": "03",
    "apr": "04", "april": "04",
    "may": "05",
    "jun": "06", "june": "06",
    "jul": "07", "july": "07",
    "aug": "08", "august": "08",
    "sep": "09", "sept": "09", "september": "09",
    "oct": "10", "october": "10",
    "nov": "11", "november": "11",
    "dec": "12", "december": "12",
}


def infer_sheet_date(sheet_name: str, file_name: str) -> str:
    """Return YYYY-MM-DD best-guess date from sheet title / file name, or empty string."""
    combined = (f"{sheet_name} {file_name}").lower()
    for word, month_num in _MONTH_MAP.items():
        if word in combined:
            # Try to find a day number right after the month word
            m = re.search(re.escape(word) + r"\s*(\d{1,2})", combined)
            day = "01"
            if m:
                try:
                    d = int(m.group(1))
                    if 1 <= d <= 31:
                        day = str(d).zfill(2)
                except ValueError:
                    pass
            return f"2025-{month_num}-{day}"
    # File "Customer Details.xlsx" (no month tagged) is the main August workbook
    if re.sub(r"\s+", "", file_name.lower()) in ("customerdetails.xlsx", "customerdetails"):
        return "2025-08-01"
    return ""


def make_source_id(order_id: str, name: str, phone: str, created_ts: str, file_name: str, sheet_name: str, row_num: int) -> str:
    raw = "|".join([order_id, name, phone, created_ts, file_name, sheet_name, str(row_num)])
    return "xlsx_" + hashlib.sha1(raw.encode("utf-8")).hexdigest()[:12]


def extract_rows_from_workbook(path: Path, batch_tag: str, unknown_acc: dict[str, dict[str, Any]]) -> tuple[list[dict[str, str]], int]:
    try:
        from openpyxl import load_workbook
    except ModuleNotFoundError as exc:
        raise SystemExit(
            "Missing dependency 'openpyxl'. Install it or run in your backend Python environment."
        ) from exc

    wb = load_workbook(path, data_only=True, read_only=True)
    all_rows: list[dict[str, str]] = []
    sheet_count = 0

    for ws in wb.worksheets:
        # Skip metadata / leads sheets that contain no order rows
        ws_name_lower = ws.title.strip().lower()
        if ws_name_lower in NON_ORDER_SHEETS:
            continue
        if re.match(r"^sheet\d+$", ws_name_lower):
            continue

        sheet_count += 1

        try:
            header_row, col_map, item_cols = find_header_row(ws)
        except ValueError:
            # Sheet doesn't have enough structure to identify order data
            sheet_count -= 1
            continue

        # Fallback: try to sniff the items column from data values
        if not item_cols:
            item_cols = infer_items_column(ws, header_row, col_map)

        empty_streak = 0

        # Pre-compute a sheet-level fallback date for rows without dates or order IDs
        sheet_fallback_date = infer_sheet_date(ws.title, path.name)

        for row_idx, row in enumerate(
            ws.iter_rows(min_row=header_row + 1, values_only=True),
            start=header_row + 1,
        ):
            values = [as_text(v) for v in row]

            def get_value(col_name: str) -> str:
                idx = col_map.get(col_name)
                if idx is None or idx >= len(values):
                    return ""
                return (values[idx] or "").strip()

            order_id = normalize_order_id(get_value("order_id"))
            name = get_value("name")
            address = get_value("address")
            phone = get_value("phone")
            amount = format_amount(get_value("amount"))

            # Empty-streak check (detects end-of-data region at bottom of sheet)
            if not any([order_id, name, address, phone, amount]):
                empty_streak += 1
                if empty_streak >= 20:
                    break
                continue
            empty_streak = 0

            # Minimum data quality: must have at least name or phone to identify customer
            if not name and not phone:
                continue

            raw_item_parts: list[str] = []
            for idx in item_cols:
                if idx < len(values):
                    value = (values[idx] or "").strip()
                    if value:
                        raw_item_parts.append(value)
            raw_items = "\n".join(raw_item_parts)

            if not raw_items:
                raw_items = get_value("items")

            tracking_id = get_value("tracking_id")
            # Clear placeholder hyperlink text
            if tracking_id.lower() in ("track", "tracking"):
                tracking_id = ""
            pincode = get_value("pincode")
            payment_type = get_value("payment_type")
            # Convert Paid (Yes/No) → prepaid / cod
            pt_lower = payment_type.lower().strip()
            if pt_lower in ("yes", "y", "paid", "1", "true"):
                payment_type = "prepaid"
            elif pt_lower in ("no", "n", "0", "false", "cod"):
                payment_type = "cod"
            source_notes = get_value("notes")
            source_id = get_value("_id")
            printed_count = get_value("printed_count")
            created_raw = get_value("created_ts")

            phone_display = get_value("phone_display")
            if not phone_display and phone:
                phone_display = format_phone_display(phone)

            normalized_items, unknown_tokens = normalize_items(raw_items, amount)
            for token in unknown_tokens:
                bucket = unknown_acc[token]
                bucket["count"] += 1
                bucket["order_ids"].add(order_id)
                bucket["customers"].add(name)

            created_ts = to_iso_created_ts(created_raw, order_id, sheet_fallback_date)
            if not source_id:
                source_id = make_source_id(order_id, name, phone, created_ts, path.name, ws.title, row_idx)

            try:
                pc_int = int(float(printed_count)) if printed_count else 0
            except ValueError:
                pc_int = 0

            provenance = (
                f"BATCH_SOURCE=XLSX;BATCH_TAG={batch_tag};FILE={path.name};"
                f"SHEET={ws.title};ROW={row_idx}"
            )
            notes = " | ".join([x for x in [source_notes, provenance] if x])

            all_rows.append(
                {
                    "name": name,
                    "address": address,
                    "pincode": pincode,
                    "phone": phone,
                    "phone_display": phone_display,
                    "payment_type": payment_type,
                    "amount": amount,
                    "order_id": order_id,
                    "tracking_id": tracking_id,
                    "items": normalized_items,
                    "notes": notes,
                    "_id": source_id,
                    "created_ts": created_ts,
                    "printed_count": str(pc_int),
                }
            )

    return all_rows, sheet_count


def write_unknown_items(path: Path, unknown_acc: dict[str, dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(
            f,
            fieldnames=["source_item_name", "occurrences", "sample_order_ids", "sample_customers"],
        )
        w.writeheader()
        for token, payload in sorted(unknown_acc.items(), key=lambda kv: kv[1]["count"], reverse=True):
            w.writerow(
                {
                    "source_item_name": token,
                    "occurrences": payload["count"],
                    "sample_order_ids": "; ".join(sorted(list(payload["order_ids"]))[:5]),
                    "sample_customers": "; ".join(sorted(list(payload["customers"]))[:5]),
                }
            )


def main() -> int:
    args = parse_args()

    input_paths = [Path(p) for p in args.input]
    for p in input_paths:
        if not p.exists():
            raise SystemExit(f"Input XLSX not found: {p}")

    unknown_acc: dict[str, dict[str, Any]] = defaultdict(
        lambda: {"count": 0, "order_ids": set(), "customers": set()}
    )

    total_rows = 0
    total_sheets = 0
    extracted_rows: list[dict[str, str]] = []
    for input_path in input_paths:
        rows, sheet_count = extract_rows_from_workbook(input_path, args.batch_tag, unknown_acc)
        total_rows += len(rows)
        total_sheets += sheet_count
        extracted_rows.extend(rows)

    out_path = Path(args.output)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    with out_path.open("w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=CANONICAL_HEADERS)
        writer.writeheader()
        writer.writerows(extracted_rows)

    if args.unknown_items_output:
        write_unknown_items(Path(args.unknown_items_output), unknown_acc)

    print(
        "XLSX extraction complete: "
        f"files={len(input_paths)}, sheets={total_sheets}, rows={total_rows}"
    )
    print(f"Raw CSV: {out_path}")
    if args.unknown_items_output:
        print(f"Unknown items CSV: {args.unknown_items_output}")

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
